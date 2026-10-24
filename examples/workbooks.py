"""Build disposable example workbooks. Neither backend is a runtime dependency."""

from __future__ import annotations

import json
import os
import shlex
import subprocess
import tempfile
from pathlib import Path


def column_letter(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def build_workbook(path, sheets, *, backend=None):
    """Write invented rows to a new workbook using openpyxl or an installed MOG."""
    backend = backend or os.environ.get("XLDIFF_TEST_BACKEND", "openpyxl")
    path = Path(path)
    if backend == "openpyxl":
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, rows in sheets.items():
            sheet = workbook.create_sheet(title=name)
            for row_index, row in enumerate(rows, start=1):
                for column_index, value in enumerate(row, start=1):
                    if value is not None:
                        sheet.cell(row=row_index, column=column_index, value=value)
        workbook.save(path)
        return path
    if backend != "mog":
        raise ValueError(f"Unknown workbook backend: {backend}")

    # Fixture assertions require exactly the requested sheets. The optional
    # local guard documents this exception instead of adding a log worksheet.
    script = """
if (typeof mogLog === "function") {
    mogLog({request: "Create an invented xldiff fixture", action: "Write exact fixture sheets", details: "Disposable test data", outcome: "Match the requested worksheet shape", agent: "xldiff tests"});
    mogLog(false, "Exact test fixtures must not contain an additional log worksheet");
}
const sheets = SHEETS_JSON;
await Excel.run(async (context) => {
    const collection = context.workbook.worksheets;
    collection.load("items/name");
    await context.sync();
    const original = Array.from(collection.items);
    const temporary = collection.add("__xldiff_fixture__");
    for (const sheet of original) sheet.delete();
    for (const [name, rows] of Object.entries(sheets)) {
        const sheet = collection.add(name);
        for (let r = 0; r < rows.length; r++) {
            for (let c = 0; c < rows[r].length; c++) {
                const value = rows[r][c];
                if (value === null) continue;
                const cell = sheet.getRangeByIndexes(r, c, 1, 1);
                if (typeof value === "string" && value.startsWith("=")) cell.formulas = [[value]];
                else cell.values = [[value]];
            }
        }
    }
    temporary.delete();
    await context.sync();
});
""".replace("SHEETS_JSON", json.dumps(sheets))
    with tempfile.TemporaryDirectory(prefix="xldiff-fixture-") as directory:
        script_path = Path(directory) / "fixture.js"
        script_path.write_text(script, encoding="utf-8")
        guard = os.environ.get("MOG_GUARD_COMMAND")
        if guard:
            subprocess.run([*shlex.split(guard), "--add", str(script_path)], check=True, capture_output=True, text=True)
        result = subprocess.run(
            [os.environ.get("MOG_BIN", "mog"), "-f", str(script_path), "-o", str(path)],
            capture_output=True, text=True, timeout=60,
        )
        if result.returncode or not path.is_file():
            raise RuntimeError(f"MOG fixture generation failed: {result.stdout}\n{result.stderr}")
    return path
